import os
import re
import json
import base64
from io import BytesIO

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from openpyxl import load_workbook


try:
    from openai import OpenAI
except ImportError:
    OpenAI = None


load_dotenv()


# =====================================================
# PAGE SETTINGS
# =====================================================
st.set_page_config(
    page_title="AI Claim Amount Calculator",
    page_icon="💗",
    layout="wide"
)


# =====================================================
# COLOURFUL PROFESSIONAL THEME
# =====================================================
st.markdown(
    """
    <style>
    /* MAIN APP BACKGROUND */
    .stApp {
        background:
            linear-gradient(rgba(23, 7, 44, 0.72), rgba(23, 7, 44, 0.72)),
            radial-gradient(circle at 10% 15%, rgba(255, 87, 178, 0.95), transparent 28%),
            radial-gradient(circle at 85% 10%, rgba(87, 117, 255, 0.90), transparent 30%),
            radial-gradient(circle at 20% 90%, rgba(20, 184, 166, 0.85), transparent 30%),
            radial-gradient(circle at 80% 85%, rgba(251, 191, 36, 0.65), transparent 26%),
            linear-gradient(135deg, #22092c 0%, #461959 35%, #872341 70%, #1f2937 100%);
        color: #ffffff !important;
    }

    /* ABSTRACT BACKGROUND PATTERN */
    .stApp::before {
        content: "";
        position: fixed;
        inset: 0;
        background-image:
            url("data:image/svg+xml,%3Csvg width='900' height='600' viewBox='0 0 900 600' xmlns='http://www.w3.org/2000/svg'%3E%3Cg fill='none' stroke='%23ffffff' stroke-opacity='0.12'%3E%3Cpath d='M0 120 C180 40 300 230 480 140 C660 50 720 200 900 120'/%3E%3Cpath d='M0 260 C160 170 320 370 500 260 C680 150 740 340 900 260'/%3E%3Cpath d='M0 420 C200 330 330 540 520 420 C700 300 760 520 900 420'/%3E%3C/g%3E%3Cg fill='%23ffffff' fill-opacity='0.08'%3E%3Ccircle cx='120' cy='110' r='5'/%3E%3Ccircle cx='780' cy='150' r='6'/%3E%3Ccircle cx='650' cy='440' r='5'/%3E%3Ccircle cx='250' cy='500' r='7'/%3E%3C/g%3E%3C/svg%3E");
        background-size: cover;
        background-position: center;
        pointer-events: none;
        z-index: -1;
    }

    /* KEEP TEXT READABLE */
    html, body, [class*="css"] {
        font-family: "Segoe UI", sans-serif;
    }

    p, span, label, div {
        color: inherit;
    }

    h1, h2, h3, h4, h5, h6 {
        color: #ffffff !important;
        font-weight: 900 !important;
    }

    /* HERO */
    .hero-box {
        padding: 42px 34px;
        border-radius: 36px;
        background:
            linear-gradient(135deg, rgba(255,255,255,0.22), rgba(255,255,255,0.08));
        border: 1px solid rgba(255,255,255,0.30);
        box-shadow:
            0 26px 80px rgba(0,0,0,0.38),
            inset 0 1px 0 rgba(255,255,255,0.30);
        backdrop-filter: blur(18px);
        margin-bottom: 30px;
        text-align: center;
    }

    .main-title {
        font-size: 56px;
        line-height: 1.08;
        font-weight: 950;
        background: linear-gradient(90deg, #ffffff, #ffe4f3, #fbcfe8, #dbeafe);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 14px;
        letter-spacing: -1px;
    }

    .subtitle {
        color: #fdf2f8 !important;
        font-size: 19px;
        line-height: 1.7;
        max-width: 1030px;
        margin: auto;
        font-weight: 650;
    }

    .badge {
        display: inline-block;
        padding: 10px 17px;
        border-radius: 999px;
        background: rgba(255,255,255,0.18);
        border: 1px solid rgba(255,255,255,0.30);
        color: #ffffff !important;
        font-weight: 850;
        margin: 7px;
        font-size: 15px;
        box-shadow: 0 8px 18px rgba(0,0,0,0.16);
    }

    /* MAIN CARDS */
    .glass-card {
        background: rgba(255, 255, 255, 0.16);
        padding: 30px;
        border-radius: 30px;
        border: 1px solid rgba(255, 255, 255, 0.28);
        box-shadow: 0 24px 70px rgba(0,0,0,0.28);
        backdrop-filter: blur(18px);
        margin-bottom: 24px;
    }

    .glass-card:hover {
        transform: translateY(-2px);
        transition: 0.25s ease;
        box-shadow: 0 34px 85px rgba(0,0,0,0.34);
    }

    .section-title {
        font-size: 30px;
        font-weight: 950;
        color: #ffffff !important;
        margin-bottom: 14px;
    }

    .note-box {
        background: rgba(255,255,255,0.18);
        border: 1px solid rgba(255,255,255,0.28);
        color: #ffffff !important;
        padding: 17px 19px;
        border-radius: 20px;
        font-size: 16px;
        line-height: 1.6;
        font-weight: 650;
        margin-bottom: 18px;
    }

    .alert-box {
        background: rgba(251, 191, 36, 0.22);
        border: 1px solid rgba(251, 191, 36, 0.45);
        color: #fff7ed !important;
        padding: 17px 19px;
        border-radius: 20px;
        font-size: 16px;
        line-height: 1.6;
        font-weight: 800;
    }

    .upload-card {
        background: rgba(255,255,255,0.20);
        border-radius: 26px;
        padding: 24px;
        border: 1px solid rgba(255,255,255,0.30);
        box-shadow: 0 18px 45px rgba(0,0,0,0.22);
        margin-top: 15px;
        margin-bottom: 18px;
        backdrop-filter: blur(15px);
    }

    .upload-title {
        font-size: 23px;
        font-weight: 950;
        color: #ffffff !important;
        margin-bottom: 8px;
    }

    .upload-text {
        font-size: 16px;
        color: #fdf2f8 !important;
        line-height: 1.7;
        font-weight: 650;
    }

    /* SUMMARY CARDS */
    .money-card {
        background: linear-gradient(135deg, #ff2d95, #8b5cf6);
        padding: 36px 24px;
        border-radius: 32px;
        color: white !important;
        text-align: center;
        box-shadow:
            0 22px 56px rgba(255, 45, 149, 0.35),
            inset 0 1px 0 rgba(255,255,255,0.35);
        min-height: 205px;
    }

    .paid-card {
        background: linear-gradient(135deg, #10b981, #06b6d4);
        padding: 36px 24px;
        border-radius: 32px;
        color: white !important;
        text-align: center;
        box-shadow:
            0 22px 56px rgba(16, 185, 129, 0.30),
            inset 0 1px 0 rgba(255,255,255,0.35);
        min-height: 205px;
    }

    .net-card {
        background: linear-gradient(135deg, #6366f1, #ec4899);
        padding: 36px 24px;
        border-radius: 32px;
        color: white !important;
        text-align: center;
        box-shadow:
            0 22px 56px rgba(99, 102, 241, 0.32),
            inset 0 1px 0 rgba(255,255,255,0.35);
        min-height: 205px;
    }

    .money-label {
        font-size: 18px;
        color: #ffffff !important;
        margin-bottom: 12px;
        font-weight: 850;
    }

    .money-value {
        font-size: 43px;
        color: #ffffff !important;
        font-weight: 950;
        letter-spacing: -1px;
    }

    .mini-card {
        background: rgba(255,255,255,0.18);
        border-radius: 28px;
        padding: 26px 18px;
        border: 1px solid rgba(255,255,255,0.30);
        box-shadow:
            0 18px 44px rgba(0,0,0,0.22),
            inset 0 1px 0 rgba(255,255,255,0.22);
        text-align: center;
        min-height: 165px;
        backdrop-filter: blur(16px);
    }

    .mini-card:hover {
        transform: translateY(-4px);
        transition: 0.25s ease;
    }

    .mini-icon {
        font-size: 38px;
        margin-bottom: 9px;
    }

    .mini-label {
        color: #fdf2f8 !important;
        font-size: 16px;
        font-weight: 850;
    }

    .mini-value {
        color: #ffffff !important;
        font-size: 36px;
        font-weight: 950;
        margin-top: 10px;
    }

    /* BUTTONS */
    .stButton > button {
        width: 100%;
        border-radius: 18px;
        height: 58px;
        font-weight: 950;
        font-size: 17px;
        background: linear-gradient(90deg, #ff2d95, #7c3aed);
        color: white !important;
        border: none;
        box-shadow: 0 16px 38px rgba(255,45,149,0.35);
        transition: 0.25s ease;
    }

    .stButton > button:hover {
        transform: translateY(-2px);
        background: linear-gradient(90deg, #ec4899, #6d28d9);
        color: white !important;
        border: none;
        box-shadow: 0 22px 54px rgba(124,58,237,0.42);
    }

    /* TABS */
    .stTabs [data-baseweb="tab-list"] {
        gap: 10px;
        background: rgba(255,255,255,0.18);
        padding: 10px;
        border-radius: 22px;
        border: 1px solid rgba(255,255,255,0.30);
        box-shadow: 0 14px 34px rgba(0,0,0,0.18);
        backdrop-filter: blur(16px);
    }

    .stTabs [data-baseweb="tab"] {
        height: 54px;
        white-space: nowrap;
        border-radius: 16px;
        padding: 0 20px;
        background: rgba(255,255,255,0.95) !important;
        color: #4c1d95 !important;
        font-weight: 950 !important;
        font-size: 16px !important;
        border: 1px solid rgba(255,255,255,0.75) !important;
    }

    .stTabs [data-baseweb="tab"] p,
    .stTabs [data-baseweb="tab"] span,
    .stTabs [data-baseweb="tab"] div {
        color: #4c1d95 !important;
        font-weight: 950 !important;
    }

    .stTabs [data-baseweb="tab"]:hover {
        background: #fdf2f8 !important;
        color: #831843 !important;
    }

    .stTabs [aria-selected="true"] {
        background: linear-gradient(90deg, #ff2d95, #7c3aed) !important;
        color: #ffffff !important;
        border: none !important;
        box-shadow: 0 10px 24px rgba(255,45,149,0.32);
    }

    .stTabs [aria-selected="true"] p,
    .stTabs [aria-selected="true"] span,
    .stTabs [aria-selected="true"] div {
        color: #ffffff !important;
    }

    /* INPUTS */
    div[data-testid="stFileUploader"] {
        background: rgba(255,255,255,0.95) !important;
        padding: 22px;
        border-radius: 24px;
        border: 2px dashed #c084fc;
        color: #3b0764 !important;
    }

    div[data-testid="stFileUploader"] * {
        color: #3b0764 !important;
        font-weight: 700;
    }

    div[data-testid="stTextArea"] textarea {
        background-color: rgba(255,255,255,0.96) !important;
        color: #3b0764 !important;
        border-radius: 20px;
        border: 1px solid #c084fc !important;
        font-size: 17px;
        line-height: 1.6;
    }

    div[data-testid="stTextArea"] textarea::placeholder {
        color: #6b7280 !important;
    }

    div[data-testid="stWidgetLabel"] label,
    div[data-testid="stWidgetLabel"] p {
        color: #ffffff !important;
        font-weight: 900 !important;
        font-size: 16px !important;
    }

    div[data-baseweb="select"] > div,
    .stSelectbox div[data-baseweb="select"] > div,
    .stTextInput input {
        background: rgba(255,255,255,0.96) !important;
        color: #3b0764 !important;
        border-radius: 14px !important;
        border: 1px solid #c084fc !important;
    }

    div[data-baseweb="select"] * {
        color: #3b0764 !important;
        font-weight: 700;
    }

    /* DATA TABLES */
    div[data-testid="stDataFrame"],
    div[data-testid="stDataEditor"] {
        background: white !important;
        border-radius: 18px;
        color: #111827 !important;
        overflow: hidden;
    }

    div[data-testid="stDataFrame"] *,
    div[data-testid="stDataEditor"] * {
        color: #111827 !important;
    }

    /* EXPANDERS */
    .streamlit-expanderHeader {
        color: #3b0764 !important;
        font-weight: 950 !important;
        background: rgba(255,255,255,0.95) !important;
        border-radius: 14px !important;
    }

    .streamlit-expanderContent {
        background: rgba(255,255,255,0.96) !important;
        color: #111827 !important;
        border-radius: 0 0 14px 14px !important;
    }

    .streamlit-expanderContent * {
        color: #111827 !important;
    }

    .stAlert {
        border-radius: 16px !important;
    }

    hr {
        border-color: rgba(255,255,255,0.32) !important;
    }

    code {
        color: #111827 !important;
        background: #fdf2f8 !important;
    }
    </style>
    """,
    unsafe_allow_html=True
)


# =====================================================
# BASIC HELPERS
# =====================================================
def is_empty_value(value):
    return value is None or str(value).strip() == "" or str(value).strip().lower() in ["nan", "none"]


def extract_currency_amount(text):
    if not isinstance(text, str):
        text = str(text)

    pattern = r"\$\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]{2})|[0-9]+(?:\.[0-9]{2}))"
    match = re.search(pattern, text)

    if match:
        return float(match.group(1).replace(",", ""))

    return None


def parse_amount_value(value):
    if is_empty_value(value):
        return None

    text = str(value).replace("$", "").replace(",", "").strip()

    try:
        return float(text)
    except ValueError:
        return None


def is_black_fill(cell):
    fill = cell.fill

    if fill is None or fill.fill_type is None:
        return False

    color = fill.fgColor

    if color is None:
        return False

    rgb = color.rgb

    if rgb:
        rgb = str(rgb).upper()

        if rgb in ["FF000000", "00000000"]:
            return True

        if len(rgb) == 8:
            rgb = rgb[2:]

        if len(rgb) == 6:
            try:
                r = int(rgb[0:2], 16)
                g = int(rgb[2:4], 16)
                b = int(rgb[4:6], 16)
                brightness = (r * 299 + g * 587 + b * 114) / 1000
                return brightness < 45
            except ValueError:
                return False

    return False


def detect_black_row(cells):
    black_cells = 0

    for cell in cells:
        if is_black_fill(cell):
            black_cells += 1

    return black_cells >= 1


# =====================================================
# ROW CLASSIFICATION
# =====================================================
def classify_claim_row(
    original_text,
    row_number,
    forced_pending_amount=None,
    forced_paid_amount=None,
    row_alert="Normal"
):
    text = "" if original_text is None else str(original_text).strip()
    clean_text = text.lower()

    pending_amount = 0.00
    paid_amount = 0.00
    include_pending = False
    include_paid = False

    if row_alert == "Blank Row":
        return {
            "row": row_number,
            "original_text": "",
            "category": "Blank Row",
            "pending_amount": 0.00,
            "paid_amount": 0.00,
            "row_alert": "Blank Row",
            "status": "Blank row found",
            "include_pending": False,
            "include_paid": False
        }

    if forced_pending_amount is not None and forced_pending_amount > 0:
        pending_amount = float(forced_pending_amount)
        include_pending = True
        category = "Pending On Amount"
        status = "Pending amount counted"

    elif forced_paid_amount is not None and forced_paid_amount > 0:
        paid_amount = float(forced_paid_amount)
        include_paid = True
        category = "Paid Amount"
        status = "Paid amount counted"

    elif "no paid" in clean_text or "no paid amnt" in clean_text or "no paid amount" in clean_text:
        category = "Pending No Paid Amount"
        status = "No paid amount"

    else:
        detected_amount = extract_currency_amount(text)

        if detected_amount is not None and "paid" in clean_text and "no paid" not in clean_text:
            paid_amount = float(detected_amount)
            include_paid = True
            category = "Paid Amount"
            status = "Paid amount counted"

        elif detected_amount is not None and ("pending on" in clean_text or "pending" in clean_text):
            pending_amount = float(detected_amount)
            include_pending = True
            category = "Pending On Amount"
            status = "Pending amount counted"

        elif clean_text == "pending" or "pending" in clean_text:
            category = "Pending Only"
            status = "Pending but no amount"

        else:
            category = "Other / Ignored"
            status = "Ignored"

    if row_alert == "Black Row":
        status = status + " | Black row found"

    return {
        "row": row_number,
        "original_text": text,
        "category": category,
        "pending_amount": pending_amount,
        "paid_amount": paid_amount,
        "row_alert": row_alert,
        "status": status,
        "include_pending": include_pending,
        "include_paid": include_paid
    }


def process_pasted_text(text):
    rows = []

    for index, line in enumerate(text.splitlines(), start=1):
        if line.strip() == "":
            rows.append(
                classify_claim_row(
                    original_text="",
                    row_number=index,
                    row_alert="Blank Row"
                )
            )
        else:
            rows.append(
                classify_claim_row(
                    original_text=line,
                    row_number=index
                )
            )

    return pd.DataFrame(rows)


# =====================================================
# FILE LOADING
# =====================================================
def load_csv_file(file_bytes):
    df = pd.read_csv(
        BytesIO(file_bytes),
        dtype=str,
        keep_default_na=False,
        skip_blank_lines=False
    )

    meta_rows = []

    for index, row in df.iterrows():
        row_values = list(row.values)
        is_blank = all(is_empty_value(value) for value in row_values)

        meta_rows.append({
            "source_index": index,
            "row_number": index + 2,
            "row_alert": "Blank Row" if is_blank else "Normal"
        })

    meta_df = pd.DataFrame(meta_rows)

    return df, meta_df


def load_excel_file(file_bytes):
    workbook = load_workbook(BytesIO(file_bytes), data_only=True)
    sheet = workbook.active

    headers = []

    for cell in sheet[1]:
        header = cell.value
        if is_empty_value(header):
            header = f"Column_{cell.column}"
        headers.append(str(header))

    records = []
    meta_rows = []

    for excel_row_number in range(2, sheet.max_row + 1):
        cells = list(sheet[excel_row_number])
        values = [cell.value for cell in cells[:len(headers)]]

        row_dict = {}

        for header, value in zip(headers, values):
            row_dict[header] = "" if value is None else value

        is_blank = all(is_empty_value(value) for value in values)
        is_black = detect_black_row(cells)

        if is_blank:
            row_alert = "Blank Row"
        elif is_black:
            row_alert = "Black Row"
        else:
            row_alert = "Normal"

        records.append(row_dict)

        meta_rows.append({
            "source_index": len(records) - 1,
            "row_number": excel_row_number,
            "row_alert": row_alert
        })

    df = pd.DataFrame(records)
    meta_df = pd.DataFrame(meta_rows)

    return df, meta_df


def process_uploaded_file(file_bytes, file_name, text_column, pending_amount_column=None, paid_amount_column=None):
    if file_name.lower().endswith(".csv"):
        source_df, meta_df = load_csv_file(file_bytes)
    else:
        source_df, meta_df = load_excel_file(file_bytes)

    rows = []

    for index, row in source_df.iterrows():
        meta_row = meta_df[meta_df["source_index"] == index].iloc[0]

        row_number = int(meta_row["row_number"])
        row_alert = meta_row["row_alert"]

        original_text = row[text_column] if text_column is not None else ""

        pending_amount = None
        paid_amount = None

        if pending_amount_column is not None:
            pending_amount = parse_amount_value(row[pending_amount_column])

        if paid_amount_column is not None:
            paid_amount = parse_amount_value(row[paid_amount_column])

        rows.append(
            classify_claim_row(
                original_text=original_text,
                row_number=row_number,
                forced_pending_amount=pending_amount,
                forced_paid_amount=paid_amount,
                row_alert=row_alert
            )
        )

    return pd.DataFrame(rows), source_df, meta_df


# =====================================================
# SCREENSHOT AI
# =====================================================
def image_to_base64(uploaded_file):
    return base64.b64encode(uploaded_file.getvalue()).decode("utf-8")


def clean_ai_json_response(text):
    text = text.strip()

    if text.startswith("```"):
        text = text.replace("```json", "").replace("```", "").strip()

    first_index = text.find("{")
    last_index = text.rfind("}")

    if first_index != -1 and last_index != -1:
        text = text[first_index:last_index + 1]

    return json.loads(text)


def process_screenshot_with_ai(uploaded_image):
    api_key = os.getenv("OPENAI_API_KEY")
    model_name = os.getenv("OPENAI_MODEL", "gpt-4o-mini")

    if not api_key:
        st.error("OPENAI_API_KEY is missing. Add it to your .env file.")
        return None

    if OpenAI is None:
        st.error("OpenAI package is not installed. Run: python -m pip install openai")
        return None

    client = OpenAI(api_key=api_key)

    image_base64 = image_to_base64(uploaded_image)

    file_extension = uploaded_image.name.split(".")[-1].lower()
    mime_type = "image/png"

    if file_extension in ["jpg", "jpeg"]:
        mime_type = "image/jpeg"
    elif file_extension == "webp":
        mime_type = "image/webp"

    prompt = """
You are an AI Claim Amount Calculator.

Read this screenshot carefully and extract rows related to claim payment status.

Return ONLY valid JSON in this exact format:

{
  "rows": [
    {
      "original_text": "full visible row text",
      "category": "Pending On Amount | Pending No Paid Amount | Paid Amount | Pending Only | Blank Row | Black Row | Other / Ignored",
      "pending_amount": 0.00,
      "paid_amount": 0.00,
      "row_alert": "Normal | Blank Row | Black Row",
      "status": "short status",
      "include_pending": true,
      "include_paid": false
    }
  ]
}

Rules:
1. If row has "pending on $amount", set category "Pending On Amount", pending_amount numeric, include_pending true.
2. If row has "pending no paid amnt", "pending no paid amount", or "no paid", set category "Pending No Paid Amount", amounts 0.
3. If row has paid amount, set category "Paid Amount", paid_amount numeric, include_paid true.
4. If row only says "pending", set category "Pending Only".
5. If there is a clearly blank row, set row_alert "Blank Row".
6. If there is a black/dark row or black-highlighted row, set row_alert "Black Row".
7. Ignore dates, IDs, phone numbers, claim numbers and reference numbers unless they are dollar amounts.
8. Return JSON only.
"""

    try:
        response = client.responses.create(
            model=model_name,
            input=[
                {
                    "role": "user",
                    "content": [
                        {"type": "input_text", "text": prompt},
                        {
                            "type": "input_image",
                            "image_url": f"data:{mime_type};base64,{image_base64}"
                        }
                    ]
                }
            ]
        )

        result_json = clean_ai_json_response(response.output_text)

        rows = []

        for index, item in enumerate(result_json.get("rows", []), start=1):
            pending_amount = parse_amount_value(item.get("pending_amount", 0)) or 0.00
            paid_amount = parse_amount_value(item.get("paid_amount", 0)) or 0.00

            rows.append({
                "row": index,
                "original_text": item.get("original_text", ""),
                "category": item.get("category", "Other / Ignored"),
                "pending_amount": pending_amount,
                "paid_amount": paid_amount,
                "row_alert": item.get("row_alert", "Normal"),
                "status": item.get("status", "Ignored"),
                "include_pending": bool(item.get("include_pending", False)),
                "include_paid": bool(item.get("include_paid", False))
            })

        return pd.DataFrame(rows)

    except Exception as error:
        st.error(f"Screenshot reading failed: {error}")
        return None


# =====================================================
# DASHBOARD FUNCTIONS
# =====================================================
def calculate_pending_total(df):
    if df.empty:
        return 0.00

    return df[df["include_pending"] == True]["pending_amount"].sum()


def calculate_paid_total(df):
    if df.empty:
        return 0.00

    return df[df["include_paid"] == True]["paid_amount"].sum()


def create_excel_download(df):
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Claim Dashboard")

    output.seek(0)
    return output


def render_dashboard(df):
    total_rows = len(df)

    pending_df = df[df["category"] == "Pending On Amount"]
    no_paid_df = df[df["category"] == "Pending No Paid Amount"]
    paid_df = df[df["category"] == "Paid Amount"]
    pending_only_df = df[df["category"] == "Pending Only"]
    blank_df = df[df["row_alert"] == "Blank Row"]
    black_df = df[df["row_alert"] == "Black Row"]
    ignored_df = df[df["category"] == "Other / Ignored"]

    pending_total = calculate_pending_total(df)
    paid_total = calculate_paid_total(df)
    net_open_amount = pending_total - paid_total

    st.markdown("---")

    t1, t2, t3 = st.columns(3)

    with t1:
        st.markdown(
            f"""
            <div class="money-card">
                <div class="money-label">Total Pending Amount</div>
                <div class="money-value">${pending_total:,.2f}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with t2:
        st.markdown(
            f"""
            <div class="paid-card">
                <div class="money-label">Total Paid Amount</div>
                <div class="money-value">${paid_total:,.2f}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with t3:
        st.markdown(
            f"""
            <div class="net-card">
                <div class="money-label">Net Open Amount</div>
                <div class="money-value">${net_open_amount:,.2f}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown("<br>", unsafe_allow_html=True)

    c1, c2, c3, c4 = st.columns(4)

    with c1:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">📌</div>
                <div class="mini-label">Total Rows / Claims</div>
                <div class="mini-value">{total_rows}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with c2:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">💗</div>
                <div class="mini-label">Pending On Amount</div>
                <div class="mini-value">{len(pending_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with c3:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">✅</div>
                <div class="mini-label">Paid Amount Rows</div>
                <div class="mini-value">{len(paid_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with c4:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">🚫</div>
                <div class="mini-label">No Paid Amount</div>
                <div class="mini-value">{len(no_paid_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown("<br>", unsafe_allow_html=True)

    c5, c6, c7, c8 = st.columns(4)

    with c5:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">⏳</div>
                <div class="mini-label">Pending Only</div>
                <div class="mini-value">{len(pending_only_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with c6:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">⬜</div>
                <div class="mini-label">Blank Rows Found</div>
                <div class="mini-value">{len(blank_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with c7:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">⬛</div>
                <div class="mini-label">Black Rows Found</div>
                <div class="mini-value">{len(black_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with c8:
        st.markdown(
            f"""
            <div class="mini-card">
                <div class="mini-icon">🧹</div>
                <div class="mini-label">Ignored / Other</div>
                <div class="mini-value">{len(ignored_df)}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    if len(blank_df) > 0 or len(black_df) > 0:
        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown(
            f"""
            <div class="alert-box">
                ⚠️ File check completed: {len(blank_df)} blank row(s) and {len(black_df)} black-highlighted row(s) found.
                Please review them in the table below.
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown("---")

    st.markdown("<div class='section-title'>Claim Amount Table</div>", unsafe_allow_html=True)

    edited_df = st.data_editor(
        df,
        use_container_width=True,
        hide_index=True,
        column_config={
            "row": st.column_config.NumberColumn("Row", width="small"),
            "original_text": st.column_config.TextColumn("Original Text", width="large"),
            "category": st.column_config.SelectboxColumn(
                "Category",
                options=[
                    "Pending On Amount",
                    "Pending No Paid Amount",
                    "Paid Amount",
                    "Pending Only",
                    "Blank Row",
                    "Black Row",
                    "Other / Ignored"
                ],
                width="medium"
            ),
            "pending_amount": st.column_config.NumberColumn("Pending Amount", format="$%.2f"),
            "paid_amount": st.column_config.NumberColumn("Paid Amount", format="$%.2f"),
            "row_alert": st.column_config.SelectboxColumn(
                "Row Alert",
                options=["Normal", "Blank Row", "Black Row"],
                width="medium"
            ),
            "status": st.column_config.TextColumn("Status", width="medium"),
            "include_pending": st.column_config.CheckboxColumn("Include Pending"),
            "include_paid": st.column_config.CheckboxColumn("Include Paid")
        }
    )

    updated_pending_total = calculate_pending_total(edited_df)
    updated_paid_total = calculate_paid_total(edited_df)
    updated_net_total = updated_pending_total - updated_paid_total

    st.markdown("<br>", unsafe_allow_html=True)

    u1, u2, u3 = st.columns(3)

    with u1:
        st.markdown(
            f"""
            <div class="money-card">
                <div class="money-label">Updated Pending Total</div>
                <div class="money-value">${updated_pending_total:,.2f}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with u2:
        st.markdown(
            f"""
            <div class="paid-card">
                <div class="money-label">Updated Paid Total</div>
                <div class="money-value">${updated_paid_total:,.2f}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with u3:
        st.markdown(
            f"""
            <div class="net-card">
                <div class="money-label">Updated Net Open Amount</div>
                <div class="money-value">${updated_net_total:,.2f}</div>
            </div>
            """,
            unsafe_allow_html=True
        )

    st.markdown("---")

    st.markdown("<div class='section-title'>Breakdown Review</div>", unsafe_allow_html=True)

    b1, b2, b3 = st.columns(3)

    with b1:
        with st.expander("💗 Pending Amount Rows", expanded=True):
            st.dataframe(
                edited_df[edited_df["category"] == "Pending On Amount"][["row", "original_text", "pending_amount", "row_alert"]],
                use_container_width=True,
                hide_index=True
            )

    with b2:
        with st.expander("✅ Paid Amount Rows", expanded=True):
            st.dataframe(
                edited_df[edited_df["category"] == "Paid Amount"][["row", "original_text", "paid_amount", "row_alert"]],
                use_container_width=True,
                hide_index=True
            )

    with b3:
        with st.expander("🚫 No Paid Amount Rows", expanded=True):
            st.dataframe(
                edited_df[edited_df["category"] == "Pending No Paid Amount"][["row", "original_text", "status", "row_alert"]],
                use_container_width=True,
                hide_index=True
            )

    b4, b5 = st.columns(2)

    with b4:
        with st.expander("⬜ Blank Rows Found", expanded=True):
            st.dataframe(
                edited_df[edited_df["row_alert"] == "Blank Row"][["row", "original_text", "status"]],
                use_container_width=True,
                hide_index=True
            )

    with b5:
        with st.expander("⬛ Black Rows Found", expanded=True):
            st.dataframe(
                edited_df[edited_df["row_alert"] == "Black Row"][["row", "original_text", "category", "status"]],
                use_container_width=True,
                hide_index=True
            )

    st.markdown("---")

    st.markdown("<div class='section-title'>Formula Breakdown</div>", unsafe_allow_html=True)

    pending_amounts = edited_df[edited_df["include_pending"] == True]["pending_amount"].tolist()
    paid_amounts = edited_df[edited_df["include_paid"] == True]["paid_amount"].tolist()

    pending_formula = " + ".join([f"${amount:,.2f}" for amount in pending_amounts])
    paid_formula = " + ".join([f"${amount:,.2f}" for amount in paid_amounts])

    with st.expander("Show full calculation formula"):
        st.write("Pending Amount Formula")
        st.code(pending_formula if pending_formula else "No pending amount selected.")

        st.write("Paid Amount Formula")
        st.code(paid_formula if paid_formula else "No paid amount selected.")

        st.success(f"Pending Total = ${updated_pending_total:,.2f}")
        st.success(f"Paid Total = ${updated_paid_total:,.2f}")
        st.info(f"Net Open Amount = ${updated_net_total:,.2f}")

    csv_file = edited_df.to_csv(index=False).encode("utf-8")
    excel_file = create_excel_download(edited_df)

    d1, d2 = st.columns(2)

    with d1:
        st.download_button(
            label="⬇️ Download CSV Report",
            data=csv_file,
            file_name="ai_claim_amount_report.csv",
            mime="text/csv",
            use_container_width=True
        )

    with d2:
        st.download_button(
            label="⬇️ Download Excel Report",
            data=excel_file,
            file_name="ai_claim_amount_report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )


# =====================================================
# HEADER
# =====================================================
st.markdown(
    """
    <div class="hero-box">
        <div class="main-title">AI Claim Amount Calculator</div>
        <div class="subtitle">
            A colourful professional claim calculator for pending, paid, and no-paid amount rows.
            Upload Excel/CSV, paste text, or attach a screenshot. It also highlights blank rows and black-coloured rows.
        </div>
        <br>
        <span class="badge">💗 Pending Amount</span>
        <span class="badge">✅ Paid Amount</span>
        <span class="badge">🚫 No Paid Amount</span>
        <span class="badge">⬜ Blank Row Check</span>
        <span class="badge">⬛ Black Row Check</span>
        <span class="badge">📤 Export Report</span>
    </div>
    """,
    unsafe_allow_html=True
)


# =====================================================
# INPUT TABS
# =====================================================
tab1, tab2, tab3 = st.tabs(
    [
        "📋 Paste Text",
        "📊 Upload Excel / CSV",
        "📸 Attach Screenshot"
    ]
)


# =====================================================
# TAB 1: PASTE TEXT
# =====================================================
with tab1:
    st.markdown("<div class='glass-card'>", unsafe_allow_html=True)

    st.markdown("<div class='section-title'>Paste Claim Data</div>", unsafe_allow_html=True)

    st.markdown(
        """
        <div class="note-box">
            Paste your claim list here. The calculator will separate pending amount, paid amount,
            no-paid amount, pending-only rows, and blank rows.
        </div>
        """,
        unsafe_allow_html=True
    )

    sample_text = """pending no paid amnt
pending on $2,866.02
paid on $500.00
pending
pending on $183.75
paid $1,000.00"""

    user_text = st.text_area(
        "Paste your claim data",
        height=320,
        placeholder=sample_text
    )

    if st.button("🚀 Generate Dashboard from Text"):
        if user_text.strip():
            st.session_state["claim_dashboard_df"] = process_pasted_text(user_text)
            st.success("Dashboard generated from pasted text.")
        else:
            st.warning("Please paste data first.")

    st.markdown("</div>", unsafe_allow_html=True)


# =====================================================
# TAB 2: EXCEL / CSV
# =====================================================
with tab2:
    st.markdown("<div class='glass-card'>", unsafe_allow_html=True)

    st.markdown("<div class='section-title'>Upload Excel or CSV File</div>", unsafe_allow_html=True)

    st.markdown(
        """
        <div class="note-box">
            Upload your Excel or CSV file. For Excel .xlsx files, the app can detect blank rows and black-coloured rows.
            Select the status/text column, pending amount column, and paid amount column if available.
        </div>
        """,
        unsafe_allow_html=True
    )

    uploaded_file = st.file_uploader(
        "Upload Excel or CSV",
        type=["xlsx", "csv"],
        key="claim_file_upload"
    )

    if uploaded_file:
        file_bytes = uploaded_file.getvalue()

        try:
            if uploaded_file.name.lower().endswith(".csv"):
                preview_df, meta_df = load_csv_file(file_bytes)
                black_row_message = "CSV files do not store row colour, so black-row detection is not available for CSV."
            else:
                preview_df, meta_df = load_excel_file(file_bytes)
                black_count_preview = len(meta_df[meta_df["row_alert"] == "Black Row"])
                blank_count_preview = len(meta_df[meta_df["row_alert"] == "Blank Row"])
                black_row_message = f"Pre-check: {blank_count_preview} blank row(s), {black_count_preview} black-coloured row(s) found."

            st.markdown(
                f"""
                <div class="upload-card">
                    <div class="upload-title">📎 File Attached Successfully</div>
                    <div class="upload-text">
                        <b>File name:</b> {uploaded_file.name}<br>
                        <b>File size:</b> {len(file_bytes) / 1024:,.1f} KB<br>
                        {black_row_message}
                    </div>
                </div>
                """,
                unsafe_allow_html=True
            )

            st.write("File Preview")
            st.dataframe(preview_df.head(25), use_container_width=True)

            column_options = list(preview_df.columns)

            text_column = st.selectbox(
                "Select status/text column",
                column_options
            )

            pending_amount_options = ["Auto detect from text"] + column_options

            pending_amount_choice = st.selectbox(
                "Select pending amount column, if separate",
                pending_amount_options
            )

            paid_amount_options = ["No separate paid amount column"] + column_options

            paid_amount_choice = st.selectbox(
                "Select paid amount column, if separate",
                paid_amount_options
            )

            if st.button("🚀 Generate Dashboard from Excel / CSV"):
                pending_amount_column = None
                paid_amount_column = None

                if pending_amount_choice != "Auto detect from text":
                    pending_amount_column = pending_amount_choice

                if paid_amount_choice != "No separate paid amount column":
                    paid_amount_column = paid_amount_choice

                result_df, source_df, source_meta = process_uploaded_file(
                    file_bytes=file_bytes,
                    file_name=uploaded_file.name,
                    text_column=text_column,
                    pending_amount_column=pending_amount_column,
                    paid_amount_column=paid_amount_column
                )

                st.session_state["claim_dashboard_df"] = result_df
                st.success("Dashboard generated from Excel / CSV.")

        except Exception as error:
            st.error(f"Could not process uploaded file: {error}")

    st.markdown("</div>", unsafe_allow_html=True)


# =====================================================
# TAB 3: SCREENSHOT
# =====================================================
with tab3:
    st.markdown("<div class='glass-card'>", unsafe_allow_html=True)

    st.markdown("<div class='section-title'>Attach Screenshot for AI Calculation</div>", unsafe_allow_html=True)

    st.markdown(
        """
        <div class="note-box">
            Attach the screenshot only. The preview is hidden so it will not take over the screen.
            The AI will read pending amount, paid amount, no-paid amount, blank rows, and black rows in the background.
        </div>
        """,
        unsafe_allow_html=True
    )

    uploaded_image = st.file_uploader(
        "Attach screenshot",
        type=["png", "jpg", "jpeg", "webp"],
        key="claim_screenshot_upload"
    )

    if uploaded_image:
        file_size_kb = len(uploaded_image.getvalue()) / 1024

        st.markdown(
            f"""
            <div class="upload-card">
                <div class="upload-title">📎 Screenshot Attached Successfully</div>
                <div class="upload-text">
                    <b>File name:</b> {uploaded_image.name}<br>
                    <b>File size:</b> {file_size_kb:,.1f} KB<br>
                    Preview is hidden. This screenshot will be used only for AI calculation.
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button("🤖 Read Screenshot and Generate Dashboard"):
            with st.spinner("AI is reading your screenshot and calculating..."):
                screenshot_df = process_screenshot_with_ai(uploaded_image)

            if screenshot_df is not None and not screenshot_df.empty:
                st.session_state["claim_dashboard_df"] = screenshot_df
                st.success("Dashboard generated from screenshot.")
            else:
                st.warning("No claim rows found from screenshot.")

    st.markdown("</div>", unsafe_allow_html=True)


# =====================================================
# SHOW DASHBOARD
# =====================================================
if "claim_dashboard_df" in st.session_state:
    render_dashboard(st.session_state["claim_dashboard_df"])